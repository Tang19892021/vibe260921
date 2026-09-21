"""
네이버 검색 결과에서 신문기사 링크를 모으고, 각 기사의 본문을 크롤링하는 프로그램.

사용법
    python news_crawler.py                      # 기본 URL(검색어: 반도체)에서 기사 10개 수집
    python news_crawler.py --limit 5            # 5개만
    python news_crawler.py --source naver       # 네이버 뉴스 링크만 (본문 구조가 일정해 가장 안정적)
    python news_crawler.py --url "https://search.naver.com/search.naver?where=news&query=AI"

필요한 패키지
    pip install requests beautifulsoup4 openpyxl     (openpyxl은 엑셀 저장에만 필요)

주의
    - 개인 학습·연구 용도로 소량만 수집하세요. 요청 사이에 --delay(기본 1초)만큼 쉽니다.
    - 사이트의 이용약관과 robots.txt를 확인하세요. 기사 저작권은 각 언론사에 있습니다.
    - 네이버 검색 페이지의 클래스 이름은 자주 바뀝니다. 그래서 클래스가 아니라 링크 주소 패턴으로 기사를 찾습니다.
"""

import argparse
import csv
import json
import re
import sys
import time
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

DEFAULT_URL = (
    "https://search.naver.com/search.naver?where=nexearch&sm=top_hty&fbm=0&ie=utf8"
    "&query=%EB%B0%98%EB%8F%84%EC%B2%B4&ackey=6zjfxkkq"
)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
}

# 네이버 뉴스 기사: https://n.news.naver.com/mnews/article/언론사코드/기사번호
NAVER_NEWS_RE = re.compile(r"^https?://n\.news\.naver\.com/(?:m)?news/article/\d+/\d+")
NAVER_HOST_RE = re.compile(r"(^|\.)(naver\.com|naver\.net|pstatic\.net|navercorp\.com)$")
# 언론사 원문 링크로 보이는 경로: /article/, /news/, /view/ 또는 /2026/09/21/ 같은 날짜 경로
PRESS_PATH_RE = re.compile(r"/(news|article|articles|view|newsview)/|/\d{4}/\d{2}/\d{2}/", re.I)

# 언론사 사이트마다 다른 본문 영역. 위에서부터 순서대로 시도한다.
BODY_SELECTORS = [
    "#dic_area", "#newsct_article", "#articeBody",           # 네이버 뉴스
    "#articleBody", "#article-view-content-div", "#articletxt", "#article_body",
    "[itemprop=articleBody]", "div.article_body", "div.article-body", "div.news_body",
    "#newsContent", "div.article_txt", "article",
]
# 본문 안에 섞여 있는 사진 설명, 광고 등 불필요한 요소
NOISE_SELECTORS = (
    "script, style, iframe, figure, figcaption, .img_desc, .end_photo_org, "
    ".ad, .ads, [class*=advert], [class*=related], .byline"
)
MIN_BODY_LEN = 200


# ---------- 요청 ----------
def get_soup(session, url, retries=2):
    """URL을 받아 BeautifulSoup 객체로 돌려준다. 일시적인 오류는 재시도한다."""
    last_error = None
    for attempt in range(retries + 1):
        try:
            res = session.get(url, timeout=10)
            res.raise_for_status()
            # 헤더에 charset이 없으면 bytes를 그대로 넘겨 bs4가 <meta charset>으로 판별하게 한다
            declared = "charset" in res.headers.get("content-type", "").lower()
            return BeautifulSoup(res.content, "html.parser", from_encoding=res.encoding if declared else None)
        except requests.RequestException as e:
            last_error = e
            time.sleep(1 + attempt)
    raise last_error


# ---------- 1단계: 검색 결과에서 기사 링크 수집 ----------
def collect_links(soup, base_url, source="all"):
    """검색 결과 페이지에서 (종류, URL) 목록을 나타난 순서대로 돌려준다. 종류는 'naver' 또는 'press'."""
    links, seen = [], set()
    for a in soup.select("a[href]"):
        url = urljoin(base_url, a["href"]).split("#")[0]
        parsed = urlparse(url)
        if parsed.scheme not in ("http", "https") or url in seen:
            continue
        if NAVER_NEWS_RE.match(url):
            kind = "naver"
        elif not NAVER_HOST_RE.search(parsed.netloc) and PRESS_PATH_RE.search(parsed.path):
            kind = "press"
        else:
            continue
        if source != "all" and source != kind:
            continue
        seen.add(url)
        links.append((kind, url))
    return links


# ---------- 2단계: 기사 페이지에서 본문 추출 ----------
def meta(soup, *names):
    """<meta property|name="..."> 의 content 값을 첫 번째로 찾은 것으로 돌려준다."""
    for name in names:
        tag = soup.find("meta", attrs={"property": name}) or soup.find("meta", attrs={"name": name})
        if tag and tag.get("content"):
            return tag["content"].strip()
    return ""


# 본문 영역 안에 텍스트로 남는 광고·이미지 버튼 문구 (줄 전체가 이 문구일 때만 제거)
NOISE_LINE_RE = re.compile(r"^(AD|광고|원본보기 아이콘|이미지 확대|사진 확대|확대보기)$")
DATE_RE = re.compile(r"(\d{4})[-./](\d{2})[-./](\d{2})[T ]+(\d{2}):(\d{2})")


def normalize_date(raw):
    """'2026-09-21T09:20:51+0900', '2026/09/21 09:15' 같은 값을 'YYYY-MM-DD HH:MM'으로 통일한다."""
    m = DATE_RE.search(raw or "")
    return "{}-{}-{} {}:{}".format(*m.groups()) if m else (raw or "")


def clean_body(tag):
    """본문 태그에서 잡음을 제거하고 줄바꿈이 정리된 텍스트로 만든다."""
    for noise in tag.select(NOISE_SELECTORS):
        noise.decompose()
    for br in tag.find_all("br"):
        br.replace_with("\n")
    lines = (re.sub(r"[ \t ]+", " ", line).strip() for line in tag.get_text().splitlines())
    return "\n".join(line for line in lines if line and not NOISE_LINE_RE.match(line))


FOOTER_START_RE = re.compile(r"저작권자|무단\s?전재|Copyright|ⓒ|©", re.I)
FOOTER_LINE_RE = re.compile(r"View English Article|영문기사 보기|^[\w.+-]+@[\w-]+(\.[\w-]+)+$", re.I)


def trim_footer(text):
    """언론사 원문 끝에 붙는 저작권 문구, 영문기사 링크, 기자 이메일 줄을 잘라낸다."""
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0 and FOOTER_START_RE.search(line):
            lines = lines[:i]
            break
    while lines and FOOTER_LINE_RE.search(lines[-1].strip()):
        lines.pop()
    return "\n".join(lines)


def guess_body(soup):
    """알려진 선택자로 못 찾았을 때: 문단(<p>) 글자 수가 가장 많은 부모 요소를 본문으로 본다."""
    scores = {}
    for p in soup.find_all("p"):
        text = p.get_text(" ", strip=True)
        if len(text) >= 30:
            entry = scores.setdefault(id(p.parent), [p.parent, 0])
            entry[1] += len(text)
    if not scores:
        return ""
    parent = max(scores.values(), key=lambda e: e[1])[0]
    return clean_body(parent)


def extract_body(soup):
    for selector in BODY_SELECTORS:
        tag = soup.select_one(selector)
        if tag:
            text = trim_footer(clean_body(tag))
            if len(text) >= MIN_BODY_LEN:
                return text
    text = trim_footer(guess_body(soup))
    return text if len(text) >= MIN_BODY_LEN else ""


def extract_article(soup, url):
    """기사 페이지 하나에서 제목·언론사·날짜·본문을 뽑는다. 본문을 못 찾으면 None."""
    body = extract_body(soup)
    if not body:
        return None

    title_tag = soup.select_one("#title_area")   # 네이버 뉴스 제목
    title = title_tag.get_text(strip=True) if title_tag else (
        meta(soup, "og:title") or (soup.title.get_text(strip=True) if soup.title else ""))

    press = meta(soup, "og:site_name", "og:article:author")
    logo = soup.select_one(".media_end_head_top_logo img")   # 네이버 뉴스 언론사 로고
    if logo:
        press = logo.get("title") or logo.get("alt") or press
    press = press or urlparse(url).netloc

    date_tag = soup.select_one(".media_end_head_info_datestamp_time")   # 네이버 뉴스 작성 시각
    date = (date_tag.get("data-date-time") or date_tag.get_text(strip=True)) if date_tag else ""
    if not date:
        date = meta(soup, "article:published_time", "og:regDate", "date")
    if not date:
        time_tag = soup.find("time", attrs={"datetime": True})
        date = time_tag["datetime"] if time_tag else ""

    return {"title": title, "press": press, "date": normalize_date(date), "url": url, "body": body}


# ---------- 저장 ----------
def save_json(results, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)


def save_csv(results, path):
    # utf-8-sig: 엑셀에서 열어도 한글이 깨지지 않는다
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["title", "press", "date", "url", "body"])
        writer.writeheader()
        writer.writerows(results)


XLSX_COLUMNS = [("제목", "title", 46), ("언론사", "press", 14), ("날짜", "date", 17), ("URL", "url", 42), ("본문", "body", 100)]
XLSX_CELL_LIMIT = 32767                                   # 엑셀 셀 하나에 넣을 수 있는 최대 글자 수
XLSX_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")   # 엑셀 파일에 넣을 수 없는 제어문자


def save_xlsx(results, path):
    """엑셀(.xlsx)로 저장한다. openpyxl이 필요하다: pip install openpyxl"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "기사"
    ws.append([label for label, _key, _width in XLSX_COLUMNS])
    for item in results:
        ws.append([XLSX_ILLEGAL_RE.sub("", str(item.get(key) or ""))[:XLSX_CELL_LIMIT] for _l, key, _w in XLSX_COLUMNS])

    header_fill = PatternFill("solid", fgColor="2563EB")
    thin = Side(style="thin", color="D0D7E2")
    for col, (_label, key, width) in enumerate(XLSX_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        head = ws.cell(row=1, column=col)
        head.font = Font(bold=True, color="FFFFFF")
        head.fill = header_fill
        head.alignment = Alignment(horizontal="center", vertical="center")
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if row == 1:
                continue
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"    # 기사 제목이 '='로 시작해도 수식으로 실행되지 않게 한다
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("title", "body"))
            if key == "url" and cell.value and len(cell.value) <= 2000:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"           # 스크롤해도 헤더 행이 보이게
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def save(results, prefix):
    save_json(results, f"{prefix}.json")
    save_csv(results, f"{prefix}.csv")
    try:
        save_xlsx(results, f"{prefix}.xlsx")
    except ImportError:
        print("[안내] openpyxl이 없어 엑셀 파일은 저장하지 않았습니다. (pip install openpyxl)")


def main():
    parser = argparse.ArgumentParser(description="네이버 검색 결과의 신문기사 본문 크롤러")
    parser.add_argument("--url", default=DEFAULT_URL, help="기사 링크를 찾을 네이버 검색 URL")
    parser.add_argument("--limit", type=int, default=10, help="수집할 기사 수 (기본 10)")
    parser.add_argument("--source", choices=["all", "naver", "press"], default="all",
                        help="all=전체, naver=네이버 뉴스만, press=언론사 원문만")
    parser.add_argument("--delay", type=float, default=1.0, help="요청 사이 대기 시간(초)")
    parser.add_argument("--out", default="news_result", help="저장 파일 이름(확장자 제외). .json / .csv / .xlsx로 저장")
    args = parser.parse_args()

    sys.stdout.reconfigure(errors="replace")   # 콘솔이 표현하지 못하는 글자가 있어도 멈추지 않게
    session = requests.Session()
    session.headers.update(HEADERS)

    print("검색 결과 페이지를 가져오는 중...")
    try:
        search_soup = get_soup(session, args.url)
    except requests.RequestException as e:
        sys.exit(f"검색 페이지를 가져오지 못했습니다: {e}")

    links = collect_links(search_soup, args.url, args.source)
    print(f"기사 링크 {len(links)}개를 찾았습니다. 최대 {args.limit}개의 본문을 수집합니다.\n")

    results, seen_titles = [], set()
    for kind, url in links:
        if len(results) >= args.limit:
            break
        time.sleep(args.delay)
        try:
            article = extract_article(get_soup(session, url), url)
        except requests.RequestException as e:
            print(f"[건너뜀] 요청 실패: {url} ({e})")
            continue
        if not article:
            print(f"[건너뜀] 본문을 찾지 못함: {url}")
            continue
        if article["title"] in seen_titles:   # 네이버 뉴스와 언론사 원문이 같은 기사인 경우
            continue
        seen_titles.add(article["title"])
        results.append(article)

        preview = article["body"].replace("\n", " ")[:120]
        print(f"[{len(results)}] {article['title']}")
        print(f"    {article['press']} | {article['date']}")
        print(f"    {article['url']}")
        print(f"    {preview}...\n")

    if not results:
        sys.exit("수집된 기사가 없습니다. 페이지 구조가 바뀌었거나 접근이 제한되었을 수 있습니다.")
    save(results, args.out)
    print(f"완료: {len(results)}개 기사를 {args.out}.json / .csv / .xlsx 로 저장했습니다.")


if __name__ == "__main__":
    main()
